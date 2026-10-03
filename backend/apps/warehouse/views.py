"""
仓库管理视图
"""
import logging
import io
from datetime import datetime

from django.db import models
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
)
from . import records as record_service
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer,
    StockInSerializer, StockInWriteSerializer,
    StockOutSerializer, StockOutWriteSerializer,
    WarningSerializer, ApprovalSerializer
)
from .periods import sign_stock_out, complete_stock_out

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 货物 / 收发记录 / 审批 ====================

def _parse_date(value):
    if not value:
        return None
    return datetime.strptime(value, '%Y-%m-%d').date()


def _apply_period_version(queryset, request):
    """按期间/版本过滤。

    - 只给 period：返回该期间当前的全部有效记录（当前版本状态），
      每条记录带 version_no 标识其落账/最后更正的版本；
    - version=N：只看第 N 版引入或更正的记录（更正审计视角）；
    - 历史版本的完整可重现明细由封账快照接口提供。
    """
    period = request.query_params.get('period')
    if not period:
        return queryset
    queryset = queryset.filter(period=period)
    version = request.query_params.get('version')
    if version:
        queryset = queryset.filter(version_no=int(version))
    return queryset


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Goods.objects.select_related(
            'variety', 'variety__category', 'variety__category__unit'
        ).all()
        keyword = request.query_params.get('keyword')
        if keyword:
            queryset = queryset.filter(
                models.Q(name__icontains=keyword) | models.Q(code__icontains=keyword)
            )
        is_active = request.query_params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=is_active == 'true')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        total = queryset.count()
        goods = queryset.order_by('-created_at')[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': GoodsSerializer(goods, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })


class StockInListView(APIView):
    """入库记录：默认当前版本（开放期为实时数据），支持按期间与版本追溯。"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related('goods', 'operator').all()
        queryset = self._apply_filters(request, queryset)
        return self._paginate(request, queryset)

    def post(self, request):
        serializer = StockInWriteSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        record = record_service.create_stock_in(
            request.user,
            goods_id=data['goods'],
            quantity=data['quantity'],
            business_date=data.get('business_date'),
            batch_no=data.get('batch_no', ''),
            supplier=data.get('supplier', ''),
            remark=data.get('remark', ''),
        )
        logger.info("User %s created stock-in %s", request.user.username, record.id)
        return success_response(data=StockInSerializer(record).data, message='入库登记成功')

    def _apply_filters(self, request, queryset):
        queryset = _apply_period_version(queryset, request)
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        start, end = _parse_date(request.query_params.get('start_date')), _parse_date(request.query_params.get('end_date'))
        if start:
            queryset = queryset.filter(stock_in_time__date__gte=start)
        if end:
            queryset = queryset.filter(stock_in_time__date__lte=end)
        return queryset

    def _paginate(self, request, queryset):
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        total = queryset.count()
        records = queryset.order_by('-stock_in_time')[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': StockInSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })


class StockInDetailView(APIView):
    """入库记录修改/删除（封账期间被服务层拒绝）"""
    permission_classes = [IsAuthenticated]

    def put(self, request, pk):
        serializer = StockInWriteSerializer(data=request.data, partial=True)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        record = record_service.update_stock_in(
            request.user, pk,
            quantity=data.get('quantity'),
            business_date=data.get('business_date'),
            batch_no=data.get('batch_no'),
            supplier=data.get('supplier'),
            remark=data.get('remark'),
        )
        return success_response(data=StockInSerializer(record).data, message='更新成功')

    def delete(self, request, pk):
        record_service.delete_stock_in(request.user, pk)
        return success_response(message='删除成功')


class StockOutListView(APIView):
    """出库记录：默认当前版本，支持按期间与版本追溯。"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockOut.objects.select_related('goods', 'operator').all()
        queryset = self._apply_filters(request, queryset)
        return self._paginate(request, queryset)

    def post(self, request):
        serializer = StockOutWriteSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        record = record_service.create_stock_out(
            request.user,
            goods_id=data['goods'],
            quantity=data['quantity'],
            receiver=data['receiver'],
            receiver_dept=data.get('receiver_dept', ''),
            business_date=data.get('business_date'),
            remark=data.get('remark', ''),
        )
        logger.info("User %s created stock-out %s", request.user.username, record.id)
        return success_response(data=StockOutSerializer(record).data, message='出库申请已提交')

    def _apply_filters(self, request, queryset):
        queryset = _apply_period_version(queryset, request)
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        return queryset

    def _paginate(self, request, queryset):
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        total = queryset.count()
        records = queryset.order_by('-created_at')[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': StockOutSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })


class StockOutDetailView(APIView):
    """出库记录修改/删除（仅待审批单可直接改，封账期间被服务层拒绝）"""
    permission_classes = [IsAuthenticated]

    def put(self, request, pk):
        serializer = StockOutWriteSerializer(data=request.data, partial=True)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        record = record_service.update_stock_out(
            request.user, pk,
            quantity=data.get('quantity'),
            receiver=data.get('receiver'),
            receiver_dept=data.get('receiver_dept'),
            business_date=data.get('business_date'),
            remark=data.get('remark'),
        )
        return success_response(data=StockOutSerializer(record).data, message='更新成功')

    def delete(self, request, pk):
        record_service.delete_stock_out(request.user, pk)
        return success_response(message='删除成功')


class StockOutSignView(APIView):
    """出库单签署（审批）：并发签署结果确定，仅首次生效"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        approved = request.data.get('approved')
        if not isinstance(approved, bool):
            return error_response(message='请给出明确的审批结果 approved=true/false')
        if not request.user.is_admin:
            return error_response(message='无权限审批', code=403)
        record = sign_stock_out(pk, request.user, approved, request.data.get('remark', ''))
        return success_response(
            data=StockOutSerializer(record).data,
            message='已通过' if approved else '已拒绝',
        )


class StockOutCompleteView(APIView):
    """已批准出库单完成发放"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        record = complete_stock_out(pk, request.user)
        return success_response(data=StockOutSerializer(record).data, message='发放完成')


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Warning.objects.select_related('goods').all()
        is_read = request.query_params.get('is_read')
        if is_read in ('true', 'false'):
            queryset = queryset.filter(is_read=is_read == 'true')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        total = queryset.count()
        warnings = queryset.order_by('-created_at')[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': WarningSerializer(warnings, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Approval.objects.select_related('stock_out', 'approver').all()
        stock_out_id = request.query_params.get('stock_out')
        if stock_out_id:
            queryset = queryset.filter(stock_out_id=stock_out_id)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        total = queryset.count()
        approvals = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': ApprovalSerializer(approvals, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })
